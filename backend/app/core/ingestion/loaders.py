"""File loaders. Everything is parsed from bytes in memory - uploads never touch disk
paths derived from user input (no path traversal / zip-slip surface)."""

import io
import re
import zipfile
from pathlib import PurePosixPath

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.ingestion.models import LoadedDocument, Section

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".vtt", ".srt"}
CODE_EXTENSIONS = {
    ".py",
    ".ts",
    ".tsx",
    ".js",
    ".jsx",
    ".java",
    ".go",
    ".rs",
    ".cs",
    ".rb",
    ".php",
    ".kt",
    ".scala",
    ".html",
    ".scss",
    ".css",
    ".sql",
    ".sh",
    ".yaml",
    ".yml",
    ".toml",
    ".json",
    ".md",
    ".dockerfile",
}
SKIP_DIRS = {
    ".git",
    "node_modules",
    "dist",
    "build",
    ".venv",
    "venv",
    "__pycache__",
    ".angular",
    ".next",
    "coverage",
    ".mypy_cache",
    ".ruff_cache",
    ".pytest_cache",
    "data",
    "static",
}
SKIP_FILES = {"package-lock.json", "uv.lock", "yarn.lock", "pnpm-lock.yaml"}
MAX_CODE_FILE_BYTES = 200_000


class UnsupportedFileError(ValueError):
    pass


def _title(name: str) -> str:
    return PurePosixPath(name).stem.replace("_", " ").replace("-", " ").strip().title()


def load_bytes(filename: str, data: bytes) -> LoadedDocument:
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        sections = [
            Section(text=text, page=i)
            for i, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        ]
        return LoadedDocument(
            source=filename, title=_title(filename), doc_type="pdf", sections=sections
        )
    if suffix == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        return LoadedDocument(
            source=filename,
            title=_title(filename),
            doc_type="docx",
            sections=[Section(text=text)],
        )
    if suffix in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
        doc_type = "markdown" if suffix in {".md", ".markdown"} else "text"
        # A markdown H1 is a far better citation label than the file name.
        h1 = re.search(r"^# (.+)$", text, re.M) if doc_type == "markdown" else None
        return LoadedDocument(
            source=filename,
            title=h1.group(1).strip()[:120] if h1 else _title(filename),
            doc_type=doc_type,
            sections=[Section(text=text)],
        )
    raise UnsupportedFileError(f"Unsupported file type: {suffix or filename}")


def is_code_path(path: str) -> bool:
    parts = PurePosixPath(path).parts
    name = PurePosixPath(path).name
    if any(p in SKIP_DIRS for p in parts) or name in SKIP_FILES:
        return False
    suffix = PurePosixPath(path.lower()).suffix
    return suffix in CODE_EXTENSIONS or name.lower() in {"dockerfile", "makefile"}


def load_code_file(path: str, data: bytes, repo: str) -> LoadedDocument | None:
    if len(data) > MAX_CODE_FILE_BYTES or b"\x00" in data[:1024]:
        return None  # skip huge / binary files
    suffix = PurePosixPath(path.lower()).suffix.lstrip(".") or PurePosixPath(path).name.lower()
    return LoadedDocument(
        source=f"{repo}/{path}",
        title=path,
        doc_type="code",
        sections=[Section(text=data.decode("utf-8", errors="replace"))],
        metadata={"path": path, "lang": suffix, "repo": repo},
    )


def load_code_zip(data: bytes, repo: str, max_files: int = 1500) -> list[LoadedDocument]:
    """Read a GitHub archive zip entirely in memory."""
    docs: list[LoadedDocument] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            # GitHub archives nest everything under "<repo>-<ref>/".
            rel = "/".join(PurePosixPath(info.filename).parts[1:])
            if not rel or not is_code_path(rel) or info.file_size > MAX_CODE_FILE_BYTES:
                continue
            doc = load_code_file(rel, archive.read(info), repo)
            if doc:
                docs.append(doc)
            if len(docs) >= max_files:
                break
    return docs
